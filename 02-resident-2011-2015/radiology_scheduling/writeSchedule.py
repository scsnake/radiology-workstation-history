from openpyxl import *
import datetime
import Tkinter

clip = Tkinter.Tk().clipboard_get()

data = str(clip).split('\n')
data = [data[i].split('\t') for i in range(0, len(data))]

wb = load_workbook(filename = 'angioSchedule.xlsx')
ws = wb.active

if len(data) == 1 & len(data[0]) == 1:
    ws['A1'] = data[0][0]
else:
    for rowData in data:
        ws.append(rowData)

if ws['A1'].value == None:
    tomorrow = datetime.date.today() + datetime.timedelta(days=1)
    ws['A1'] = str(tomorrow.year) + '/' + '{:02}'.format(tomorrow.month) + '/' + '{:02}'.format(tomorrow.day)

wb.save(str(ws['A1'].value).translate(None, '/') + '.xlsx')
